"""Generate api-spec-v2.xlsx with 4 sheets:
  1. IAM 系统接口 (IAM admin endpoints)
  2. 权限点总表 (permission_groups listing across all apps)
  3. 群组 × 应用 概览 (what each Keycloak group can do)
  4. 业务接口权限矩阵 (per-endpoint × group matrix, KB + Rubik + Memory)

Source of truth for the data is da-cluster/images/keycloak-init/init-keycloak.py.
Keep this file in sync whenever a permission_group is added / changed.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

hf = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
hfill = PatternFill("solid", fgColor="2F5496")
bf = Font(name="微软雅黑", size=10)
bf_bold = Font(name="微软雅黑", size=10, bold=True)
border = Border(*[Side("thin", color="BFBFBF")] * 4)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)

method_fill = {
    "GET":    PatternFill("solid", fgColor="C6EFCE"),
    "POST":   PatternFill("solid", fgColor="BDD7EE"),
    "PUT":    PatternFill("solid", fgColor="FCE4D6"),
    "DELETE": PatternFill("solid", fgColor="FFC7CE"),
    "ANY":    PatternFill("solid", fgColor="E7E6E6"),
}

YES = "✓"      # allow (path-level pass, no further resource check — or pass-through)
RES = "⚡"     # path-level pass, resource-level ACL takes over (owner/contrib/viewer)
NO  = "✗"      # path-level deny


def _header(ws, headers):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = hf; cell.fill = hfill; cell.alignment = center; cell.border = border


def _row(ws, r, values, method_col=None):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = bf; cell.alignment = wrap if c > 3 else center; cell.border = border
        if method_col and c == method_col:
            cell.fill = method_fill.get(v, PatternFill())


# ============================================================================
# Sheet 1: IAM System APIs
# ============================================================================
IAM_HEADERS = ["模块", "接口名称", "Method", "路径", "调用方", "说明", "请求体/参数", "响应", "备注"]

IAM_APIS = [
    # --- 用户管理 ---
    ("用户管理", "用户列表", "GET", "/api/v1/{realm}/users", "前端", "支持搜索/分页/按组过滤，返回 account_type + groups", "search, group_id, first, max", "List[UserListResponse]", "account_type: internal/federated"),
    ("用户管理", "用户详情", "GET", "/api/v1/{realm}/users/{user_id}/details", "前端", "用户信息 + 所属组 + 权限（通过所在 Keycloak 组绑定的 permission_groups 聚合展开的路径）", "—", "UserDetailResponse(permissions: [{permission_group_name, path_prefix, method, required_groups, ...}])", "按应用分组展示权限"),
    ("用户管理", "创建用户", "POST", "/api/v1/{realm}/users", "前端/测试", "创建内部用户，可选绑组 + 选择是否临时密码", '{"username","password","email","groups":[gid], "temporary_password":bool=true}', "UserListResponse (201)", "temporary_password=false 用于服务/测试账号"),
    ("用户管理", "修改用户", "PUT", "/api/v1/{realm}/users/{user_id}", "前端", "修改基本信息", '{"firstName","lastName","email","enabled"}', "UserListResponse", ""),
    ("用户管理", "删除用户", "DELETE", "/api/v1/{realm}/users/{user_id}", "前端", "删除单个用户", "—", "204", "联邦用户删后IdP再登会重建"),
    ("用户管理", "批量删除", "POST", "/api/v1/{realm}/users/batch-delete", "前端", "批量删除，返回成功/失败统计", '{"user_ids":["uuid1","uuid2"]}', "BatchOperationResponse", ""),
    ("用户管理", "重置密码", "PUT", "/api/v1/{realm}/users/{user_id}/password", "前端", "重置密码(temporary=true)，联邦用户返回400", '{"password":"..."}', "204", "联邦用户无法重置"),
    ("用户管理", "添加用户到组", "PUT", "/api/v1/{realm}/users/{user_id}/groups/{group_id}", "前端", "将用户加入指定组", "—", "204", ""),
    ("用户管理", "移除用户出组", "DELETE", "/api/v1/{realm}/users/{user_id}/groups/{group_id}", "前端", "将用户从组中移除", "—", "204", ""),
    ("用户管理", "用户可选组", "GET", "/api/v1/{realm}/users/{user_id}/available-groups", "前端", "所有组+joined标记，供前端勾选", "—", '[{"id","name","joined":bool,"source"}]', ""),
    ("用户管理", "CSV导入模板", "GET", "/api/v1/{realm}/users/import-template", "前端", "下载CSV模板", "—", "text/csv attachment", "username,password,email,..."),
    ("用户管理", "批量导入", "POST", "/api/v1/{realm}/users/batch-import", "前端", "上传CSV批量创建用户", "multipart/form-data file", "BatchOperationResponse", "UTF-8/BOM均支持"),

    # --- 用户组管理 ---
    ("用户组管理", "用户组列表", "GET", "/api/v1/{realm}/groups", "前端", "支持搜索/分页，含source/member_count", "search, first, max", "List[GroupListResponse]", "source: preset/app-preset/custom"),
    ("用户组管理", "创建用户组", "POST", "/api/v1/{realm}/groups", "前端", "创建组，可选绑用户和角色", '{"name","users":[],"roles":[]}', "GroupResponse (201)", ""),
    ("用户组管理", "用户组详情", "GET", "/api/v1/{realm}/groups/{group_id}", "前端", "成员 + 角色 + 权限（通过绑定的 permission_groups 展开的路径）", "—", "GroupDetailResponse(permissions)", "按app分组展示权限"),
    ("用户组管理", "修改用户组", "PUT", "/api/v1/{realm}/groups/{group_id}", "前端", "修改名称 + 全量同步成员/角色", '{"name","users":[],"roles":[]}', "204", "users/roles 全量覆盖"),
    ("用户组管理", "删除用户组", "DELETE", "/api/v1/{realm}/groups/{group_id}", "前端", "删除自定义组，预置组返回400", "—", "204", "preset 组不可删"),
    ("用户组管理", "批量添加成员", "POST", "/api/v1/{realm}/groups/{group_id}/members/batch-add", "前端", "批量将用户添加到组", '{"user_ids":["uid1","uid2"]}', "BatchOperationResponse", ""),
    ("用户组管理", "批量移除成员", "POST", "/api/v1/{realm}/groups/{group_id}/members/batch-remove", "前端", "批量从组中移除用户", '{"user_ids":["uid1"]}', "BatchOperationResponse", ""),
    ("用户组管理", "设置组权限", "PUT", "/api/v1/{realm}/groups/{group_id}/permissions", "前端", "全量替换组绑定的 permission_groups", '{"permission_group_ids":[1,3,5]}（兼容字段 rule_ids）', '{"group_id","group_name","permission_groups":[{id, name, description, app_name}]}', "写入 permission_group_bindings 表"),

    # --- 权限管理 ---
    ("权限管理", "权限列表 (按应用分组)", "GET", "/api/v1/{realm}/permissions", "前端", "所有 permission_groups 按 app 聚合，含包含路径和绑定的 Keycloak 组", "—", '[{"app_name","app_display_name","permission_groups":[{"id","name","description","paths":[{path_prefix,method}], "bound_groups":[]}]}]', "供前端权限矩阵/勾选使用"),

    # --- 应用管理 ---
    ("应用管理", "应用列表", "GET", "/api/v1/apps", "前端/内部", "所有注册的应用", "—", "List[AppResponse]", "含 resource_patterns"),
    ("应用管理", "注册应用", "POST", "/api/v1/apps", "内部(管理员)", "注册新应用 + 资源模式 + 自动建 {app}-admins 组", '{"app_name","path_prefix","resource_patterns":[]}', "AppResponse (201)", "通常由 init 脚本预置"),
    ("应用管理", "应用详情", "GET", "/api/v1/apps/{app_name}", "前端/内部", "单个应用详情", "—", "AppResponse", ""),
    ("应用管理", "修改应用", "PUT", "/api/v1/apps/{app_name}", "内部(管理员)", "修改应用信息/License开关", '{"display_name","enabled",...}', "AppResponse", ""),
    ("应用管理", "删除应用", "DELETE", "/api/v1/apps/{app_name}", "内部(管理员)", "下线应用", "—", "204", ""),

    # --- API Key ---
    ("API Key", "创建Key", "POST", "/api/v1/{realm}/api-keys", "前端", "创建API Key，明文只返回一次", '{"name","scope","expires_at"}', "ApiKeyCreateResponse (201)", "含明文 api_key 字段"),
    ("API Key", "Key列表", "GET", "/api/v1/{realm}/api-keys", "前端", "列出所有Key（只显示前缀）", "—", "List[ApiKeyResponse]", "无明文"),
    ("API Key", "Key详情", "GET", "/api/v1/{realm}/api-keys/{key_id}", "前端", "单个Key详情", "—", "ApiKeyResponse", "无明文"),
    ("API Key", "修改Key", "PUT", "/api/v1/{realm}/api-keys/{key_id}", "前端", "修改Key信息", '{"name","scope","enabled"}', "ApiKeyResponse", ""),
    ("API Key", "删除Key", "DELETE", "/api/v1/{realm}/api-keys/{key_id}", "前端", "删除Key", "—", "204", ""),
    ("API Key", "轮换Key", "POST", "/api/v1/{realm}/api-keys/{key_id}/rotate", "前端", "轮换Key，subject_id 不变", "—", "ApiKeyCreateResponse", "旧Key立即失效"),

    # --- 资源ACL ---
    ("资源ACL", "添加权限 (分享)", "POST", "/acl/v1/resources/{resource_id}/permissions", "前端/应用", "给资源添加权限（分享）；内部校验调用者是 owner", '{"app_name","resource_type","subject_type","subject_id","permission"}', "PermissionResponse (201)", "resource-sync 提供；必须带 app_name + resource_type"),
    ("资源ACL", "查询权限", "GET", "/acl/v1/resources/{resource_id}/permissions?app_name=..&resource_type=..", "前端/应用", "查看资源的权限成员列表", "query: app_name, resource_type", '{"permissions":[PermissionResponse], "count"}', ""),
    ("资源ACL", "修改权限", "PUT", "/acl/v1/resources/{resource_id}/permissions/{acl_id}", "前端/应用", "修改权限等级；内部校验调用者是 owner", '{"app_name","resource_type","permission":"contributor"}', "PermissionResponse", ""),
    ("资源ACL", "删除权限", "DELETE", "/acl/v1/resources/{resource_id}/permissions/{acl_id}", "前端/应用", "回收权限；内部校验调用者是 owner", "—", "204", ""),

    # --- IdP/SAML 管理 ---
    ("IdP管理", "导入SAML元数据", "POST", "/api/v1/{realm}/idp/saml/import", "前端", "上传SAML元数据XML，解析IdP配置", "multipart/form-data file (XML)", "SAMLMetadataImportResponse", "返回解析出的 config 字段"),
    ("IdP管理", "创建IdP实例", "POST", "/api/v1/{realm}/idp/saml/instances", "前端", "创建SAML IdP实例（每realm仅一个）", '{"displayName","enabled","trustEmail","config":{}}', "IDPInstanceResponse (201)", "alias固定为 da-saml-idp"),
    ("IdP管理", "修改IdP实例", "PUT", "/api/v1/{realm}/idp/saml/instances", "前端", "修改SAML IdP配置", '{"displayName","enabled","trustEmail","config":{}}', "IDPInstanceResponse", ""),
    ("IdP管理", "IdP实例列表", "GET", "/api/v1/{realm}/idp/saml/instances", "前端", "列出 realm 下所有 IdP 实例", "—", "List[IDPInstanceResponse]", ""),
    ("IdP管理", "删除IdP实例", "DELETE", "/api/v1/{realm}/idp/saml/instances/{alias}", "前端", "删除指定IdP实例", "—", "204", ""),
    ("IdP管理", "Mapper列表", "GET", "/api/v1/{realm}/idp/saml/instances/{alias}/mappers", "前端", "获取 IdP 的属性映射列表（saml-user-attribute-idp-mapper）", "—", "List[IdPMapperResponse]", "仅返回属性映射类型"),
    ("IdP管理", "创建Mapper", "POST", "/api/v1/{realm}/idp/saml/instances/{alias}/mappers", "前端", "创建属性映射：SAML 属性→Keycloak 属性", '{"name","attributeKey","attributeValue","friendlyName"}', "IdPMapperResponse (201)", "固定 saml-user-attribute-idp-mapper"),
    ("IdP管理", "修改Mapper", "PUT", "/api/v1/{realm}/idp/saml/instances/{alias}/mappers/{mapper_id}", "前端", "修改属性映射", '{"name","attributeKey","attributeValue","friendlyName"}', "204", "部分更新"),
    ("IdP管理", "删除Mapper", "DELETE", "/api/v1/{realm}/idp/saml/instances/{alias}/mappers/{mapper_id}", "前端", "删除属性映射", "—", "204", ""),
    ("IdP管理", "GroupMapper列表", "GET", "/api/v1/{realm}/idp/saml/instances/{alias}/group-mappers", "前端", "获取条件化自动加组规则列表（saml-advanced-group-idp-mapper）", "—", "List[IdPGroupMapperResponse]", "每条规则：属性条件(AND)+目标组"),
    ("IdP管理", "创建GroupMapper", "POST", "/api/v1/{realm}/idp/saml/instances/{alias}/group-mappers", "前端", "创建条件化自动加组规则：SAML 属性命中时自动加入指定组", '{"name","conditions":[{"attribute","value"}],"group":"/xxx-admins","regex":false}', "IdPGroupMapperResponse (201)", "多条件AND，group必须/开头"),
    ("IdP管理", "修改GroupMapper", "PUT", "/api/v1/{realm}/idp/saml/instances/{alias}/group-mappers/{mapper_id}", "前端", "修改条件化加组规则", '{"name","conditions","group","regex"} 全可选', "204", "部分更新"),
    ("IdP管理", "删除GroupMapper", "DELETE", "/api/v1/{realm}/idp/saml/instances/{alias}/group-mappers/{mapper_id}", "前端", "删除条件化加组规则", "—", "204", ""),

    # --- Token ---
    ("Token", "授权码换Token", "POST", "/api/v1/{realm}/token/exchange", "前端", "OIDC 授权码换取 access_token", '{"code","redirect_uri","client_id","client_secret"}', "TokenExchangeResponse", "client_secret 可选"),
]


# ============================================================================
# Sheet 2: Permission Groups (business perm points)
# ============================================================================
PG_HEADERS = ["ID", "应用", "permission_group 名", "描述", "路径", "method", "绑定的 Keycloak 组"]

# Mirrors init-keycloak.py; keep these two lists in sync.
PERM_GROUPS = [
    # (app_name, name, description, [(path, method|ANY)], [kc_groups])
    # === 平台级 (app_name='') ===
    ("", "iam_admin",   "IAM 管理 API 全权",
        [("/api/v1/", "ANY")], ["admins"]),
    ("", "acl_access",  "ACL 分享 API（admins + all-users，endpoint 内部再做 owner-only 校验）",
        [("/acl/v1/", "ANY")], ["admins", "all-users"]),

    # === KB ===
    ("knowledgebase", "kb_admin_full", "KB 全应用访问",
        [("/kb/", "ANY")], ["admins", "kb-admins"]),
    ("knowledgebase", "kb_browse_knowledge_bases", "KB 列表/单查/子路径读（含 mappings/files）",
        [("/kb/knowledge_bases", "GET")], ["all-users"]),
    ("knowledgebase", "kb_browse_conversations", "会话列表/查询",
        [("/kb/conversations", "GET")], ["all-users"]),
    ("knowledgebase", "kb_browse_prompts", "提示词只读（QA 场景要用）",
        [("/kb/prompts", "GET")], ["all-users"]),
    ("knowledgebase", "kb_file_download", "文件下载（父 KB 继承）",
        [("/kb/knowledge_bases/files/download", "GET")], ["all-users"]),
    ("knowledgebase", "kb_image_download", "图片访问（父 KB 继承）",
        [("/kb/conversations/images/download", "GET")], ["all-users"]),
    ("knowledgebase", "kb_create", "KB 创建",
        [("/kb/knowledge_bases/add", "POST")], ["all-users"]),
    ("knowledgebase", "kb_modify", "KB 修改（ACL contributor）",
        [("/kb/knowledge_bases/modify", "POST")], ["all-users"]),
    ("knowledgebase", "kb_remove", "KB 删除（ACL owner）",
        [("/kb/knowledge_bases/remove", "POST")], ["all-users"]),
    ("knowledgebase", "kb_mapping_write", "目录映射写（/add, /remove）",
        [("/kb/knowledge_bases/mappings/", "POST")], ["all-users"]),
    ("knowledgebase", "kb_file_upload", "文件上传",
        [("/kb/knowledge_bases/files/upload", "POST")], ["all-users"]),
    ("knowledgebase", "kb_file_delete", "文件删除",
        [("/kb/knowledge_bases/files/remove", "POST")], ["all-users"]),
    ("knowledgebase", "kb_conv_start", "发起问答",
        [("/kb/conversations/start", "POST")], ["all-users"]),
    ("knowledgebase", "kb_conv_stop", "停止问答",
        [("/kb/conversations/stop", "GET")], ["all-users"]),
    ("knowledgebase", "kb_conv_remove", "删除会话",
        [("/kb/conversations/remove", "POST")], ["all-users"]),
    ("knowledgebase", "kb_image_generate", "生成图片链接",
        [("/kb/conversations/images/generate", "POST")], ["all-users"]),
    ("knowledgebase", "kb_conv_query", "历史对话 batch/single 查询",
        [("/kb/conversations/query/", "GET")], ["all-users"]),
    ("knowledgebase", "kb_retrieval", "检索融合搜索",
        [("/kb/retrieval/fusion_search", "POST")], ["all-users"]),
    ("knowledgebase", "kb_filesystem_manage", "文件系统管理（容器级，非 KB 级）",
        [("/kb/knowledge_bases/files/filesystem", "ANY")], ["kb-admins"]),
    ("knowledgebase", "kb_model_config_manage", "模型配置管理",
        [("/kb/models/config", "ANY")], ["kb-admins"]),
    ("knowledgebase", "kb_prompt_manage", "提示词增删改",
        [("/kb/prompts", "POST")], ["kb-admins"]),
    ("knowledgebase", "kb_jargon_group_manage", "黑话库管理",
        [("/kb/jargon_groups", "ANY")], ["kb-admins"]),
    ("knowledgebase", "kb_jargon_manage", "黑话条目管理",
        [("/kb/jargons", "ANY")], ["kb-admins"]),

    # === Rubik ===
    ("rubik", "rubik_admin_full", "Rubik 全应用访问",
        [("/rubik/", "ANY")], ["admins", "rubik-admins"]),
    ("rubik", "rubik_db_import", "option1: 导入/删除数据库",
        [("/rubik/api/databases", "POST"), ("/rubik/api/databases/", "DELETE")], ["all-users"]),
    ("rubik", "rubik_db_view", "option2: 查看数据库",
        [("/rubik/api/databases", "GET"), ("/rubik/api/databases/", "GET"), ("/rubik/api/databases/", "POST")], ["all-users"]),
    ("rubik", "rubik_db_data_dir", "option2 敏感：数据目录",
        [("/rubik/api/databases/config/data-dir", "GET")], ["rubik-admins"]),
    ("rubik", "rubik_kb_build", "option3: 知识库构建",
        [("/rubik/api/databases/", "POST"), ("/rubik/api/databases/", "GET")], ["all-users"]),
    ("rubik", "rubik_knowledge_edit", "option4: 知识增删",
        [("/rubik/api/databases/", "PUT"), ("/rubik/api/databases/", "DELETE"), ("/rubik/api/databases/", "POST")], ["all-users"]),
    ("rubik", "rubik_knowledge_special", "option4 敏感：跨库特殊知识",
        [("/rubik/api/databases/knowledge/special", "POST")], ["rubik-admins"]),
    ("rubik", "rubik_knowledge_view", "option5: 知识查看",
        [("/rubik/api/databases/", "GET"), ("/rubik/api/databases/", "POST")], ["all-users"]),
    ("rubik", "rubik_skill_manage", "option6: 技能管理",
        [("/rubik/api/databases/", "POST"), ("/rubik/api/databases/", "PUT")], ["all-users"]),
    ("rubik", "rubik_metadata", "option7: 数据库元数据管理",
        [("/rubik/api/metadata/", "GET"), ("/rubik/api/metadata/", "PUT")], ["all-users"]),
    ("rubik", "rubik_query", "option8: 问数/会话",
        [("/rubik/api/query", "POST"), ("/rubik/api/sessions", "GET"), ("/rubik/api/sessions", "POST"),
         ("/rubik/api/sessions/", "GET"), ("/rubik/api/sessions/", "POST"), ("/rubik/api/sessions/", "DELETE")],
        ["all-users"]),
    ("rubik", "rubik_config_view", "option9 读：配置查询",
        [("/rubik/api/config", "GET"), ("/rubik/api/config/", "GET")], ["all-users"]),
    ("rubik", "rubik_config_manage", "option9 写：配置管理",
        [("/rubik/api/config/", "PUT"), ("/rubik/api/config/", "POST"), ("/rubik/api/config/", "DELETE")],
        ["rubik-admins"]),
    ("rubik", "rubik_dashboard_view", "option10: 查看仪表盘",
        [("/rubik/api/dashboards", "GET"), ("/rubik/api/dashboards/", "GET")], ["all-users"]),
    ("rubik", "rubik_dashboard_edit", "option11: 增删仪表盘",
        [("/rubik/api/dashboards", "POST"), ("/rubik/api/dashboards/", "PUT"), ("/rubik/api/dashboards/", "DELETE")],
        ["all-users"]),
    ("rubik", "rubik_db_refresh", "default: 数据库刷新",
        [("/rubik/api/refresh/", "GET"), ("/rubik/api/refresh/", "POST")], ["all-users"]),

    # === Memory ===
    ("memory", "memory_admin_full", "Memory 全应用访问",
        [("/memory/", "ANY")], ["admins", "memory-admins"]),
    ("memory", "memory_health", "数据面：健康检查",
        [("/memory/api/v1/health", "GET")], ["all-users"]),
    ("memory", "memory_data_rw", "数据面：记忆增删改查（add/query/update/delete）",
        [("/memory/api/v1/memory/", "ANY")], ["all-users"]),
    ("memory", "memory_system", "数据面：系统级（故障恢复）",
        [("/memory/api/v1/system/", "ANY")], ["all-users"]),
    ("memory", "memory_tenant_manage", "管理面：租户/实例/用户记忆管理",
        [("/memory/api/v1/tenants", "ANY")], ["memory-admins"]),
    ("memory", "memory_template_manage", "管理面：模板管理",
        [("/memory/api/v1/templates", "ANY")], ["memory-admins"]),
]


# ============================================================================
# Sheet 3: 群组 × 应用 概览
# ============================================================================
OVERVIEW_HEADERS = ["Keycloak 组", "说明", "IAM 管理", "ACL 分享", "KB", "Rubik", "Memory"]

OVERVIEW_ROWS = [
    ("admins",         "平台超级管理员", "✓ 全权", "✓ 全权",              "✓ 全权", "✓ 全权", "✓ 全权"),
    ("kb-admins",      "KB 应用管理员",  "✗",      "⚡ owner",              "✓ 全权", "✗",      "✗"),
    ("rubik-admins",   "Rubik 应用管理员","✗",      "⚡ owner",              "✗",      "✓ 全权", "✗"),
    ("memory-admins",  "Memory 应用管理员","✗",    "⚡ owner",              "✗",      "✗",      "✓ 全权"),
    ("all-users",      "默认组（所有登录用户）","✗","⚡ owner (分享自己拥有的资源)",
        "细粒度读+写（资源级 ACL 再过滤）",
        "细粒度读+写（部分只给 rubik-admins）",
        "数据面全开；管理面 ✗"),
]


# ============================================================================
# Sheet 4: 业务接口权限矩阵 (endpoint × group)
# ============================================================================
ENDPOINT_HEADERS = ["应用", "功能点 / permission_group", "Method", "API 路径",
                    "admins", "kb-admins", "rubik-admins", "memory-admins", "all-users", "备注"]

# Shorthand for perm signals in rows; computed columns match OVERVIEW_ROWS order:
#   admins, kb-admins, rubik-admins, memory-admins, all-users
ENDPOINTS = [
    # ========= KB =========
    # kb_browse_knowledge_bases (all-users GET /kb/knowledge_bases)
    ("KB", "kb_browse_knowledge_bases", "GET",  "/kb/knowledge_bases/page",              YES, YES, NO,  NO,  RES, "X-Allowed-Ids 过滤"),
    ("KB", "kb_browse_knowledge_bases", "GET",  "/kb/knowledge_bases/count",             YES, YES, NO,  NO,  RES, "X-Allowed-Ids 过滤"),
    ("KB", "kb_browse_knowledge_bases", "GET",  "/kb/knowledge_bases",                   YES, YES, NO,  NO,  RES, "?KDSID=... viewer"),
    ("KB", "kb_browse_knowledge_bases", "GET",  "/kb/knowledge_bases/mappings",          YES, YES, NO,  NO,  RES, "viewer on kb"),
    ("KB", "kb_browse_knowledge_bases", "GET",  "/kb/knowledge_bases/mappings/count",    YES, YES, NO,  NO,  RES, ""),
    ("KB", "kb_browse_knowledge_bases", "GET",  "/kb/knowledge_bases/files",             YES, YES, NO,  NO,  RES, "viewer on kb"),
    ("KB", "kb_browse_knowledge_bases", "GET",  "/kb/knowledge_bases/files/history",     YES, YES, NO,  NO,  RES, ""),
    ("KB", "kb_browse_knowledge_bases", "GET",  "/kb/knowledge_bases/files/count",       YES, YES, NO,  NO,  RES, ""),
    ("KB", "kb_browse_knowledge_bases", "GET",  "/kb/knowledge_bases/{kb_name}/jargon",  YES, YES, NO,  NO,  RES, ""),
    ("KB", "kb_file_download",          "GET",  "/kb/knowledge_bases/files/download",    YES, YES, NO,  NO,  RES, "viewer 继承父 KB"),
    ("KB", "kb_browse_prompts",         "GET",  "/kb/prompts",                           YES, YES, NO,  NO,  YES, ""),
    ("KB", "kb_browse_prompts",         "GET",  "/kb/prompts/menu",                      YES, YES, NO,  NO,  YES, ""),
    ("KB", "kb_browse_conversations",   "GET",  "/kb/conversations/query/batch",         YES, YES, NO,  NO,  RES, "只看自己的"),
    ("KB", "kb_browse_conversations",   "GET",  "/kb/conversations/query/single",        YES, YES, NO,  NO,  RES, ""),
    ("KB", "kb_image_download",         "GET",  "/kb/conversations/images/download",     YES, YES, NO,  NO,  RES, "viewer 继承父 KB"),
    # Writes (all-users 精确 POST)
    ("KB", "kb_create",                 "POST", "/kb/knowledge_bases/add",               YES, YES, NO,  NO,  YES, "创建者成 owner"),
    ("KB", "kb_modify",                 "POST", "/kb/knowledge_bases/modify",            YES, YES, NO,  NO,  RES, "contributor"),
    ("KB", "kb_remove",                 "POST", "/kb/knowledge_bases/remove",            YES, YES, NO,  NO,  RES, "owner, ACL 级联删"),
    ("KB", "kb_mapping_write",          "POST", "/kb/knowledge_bases/mappings/add",      YES, YES, NO,  NO,  RES, "contributor"),
    ("KB", "kb_mapping_write",          "POST", "/kb/knowledge_bases/mappings/remove",   YES, YES, NO,  NO,  RES, "contributor"),
    ("KB", "kb_file_upload",            "POST", "/kb/knowledge_bases/files/upload",      YES, YES, NO,  NO,  RES, "contributor"),
    ("KB", "kb_file_delete",            "POST", "/kb/knowledge_bases/files/remove",      YES, YES, NO,  NO,  RES, "contributor"),
    ("KB", "kb_conv_start",             "POST", "/kb/conversations/start",               YES, YES, NO,  NO,  YES, "创建 conversation"),
    ("KB", "kb_conv_stop",              "GET",  "/kb/conversations/stop",                YES, YES, NO,  NO,  RES, "owner"),
    ("KB", "kb_conv_remove",            "POST", "/kb/conversations/remove",              YES, YES, NO,  NO,  RES, "owner"),
    ("KB", "kb_image_generate",         "POST", "/kb/conversations/images/generate",     YES, YES, NO,  NO,  RES, "viewer"),
    ("KB", "kb_retrieval",              "POST", "/kb/retrieval/fusion_search",           YES, YES, NO,  NO,  RES, "viewer"),
    # admin-only
    ("KB", "kb_filesystem_manage",      "GET",  "/kb/knowledge_bases/files/filesystem",  YES, YES, NO,  NO,  NO,  "容器级"),
    ("KB", "kb_filesystem_manage",      "POST", "/kb/knowledge_bases/files/filesystem/add",    YES, YES, NO, NO, NO, ""),
    ("KB", "kb_filesystem_manage",      "POST", "/kb/knowledge_bases/files/filesystem/remove", YES, YES, NO, NO, NO, ""),
    ("KB", "kb_model_config_manage",    "GET",  "/kb/models/config",                     YES, YES, NO,  NO,  NO,  ""),
    ("KB", "kb_model_config_manage",    "GET",  "/kb/models/config/set",                 YES, YES, NO,  NO,  NO,  ""),
    ("KB", "kb_model_config_manage",    "POST", "/kb/models/config/add",                 YES, YES, NO,  NO,  NO,  ""),
    ("KB", "kb_model_config_manage",    "POST", "/kb/models/config/modify",              YES, YES, NO,  NO,  NO,  ""),
    ("KB", "kb_model_config_manage",    "POST", "/kb/models/config/remove",              YES, YES, NO,  NO,  NO,  ""),
    ("KB", "kb_prompt_manage",          "POST", "/kb/prompts/add",                       YES, YES, NO,  NO,  NO,  ""),
    ("KB", "kb_prompt_manage",          "POST", "/kb/prompts/modify",                    YES, YES, NO,  NO,  NO,  ""),
    ("KB", "kb_prompt_manage",          "POST", "/kb/prompts/remove",                    YES, YES, NO,  NO,  NO,  ""),
    ("KB", "kb_prompt_manage",          "POST", "/kb/prompts/page",                      YES, YES, NO,  NO,  NO,  "分页（POST 归为 kb-admins，注意）"),
    ("KB", "kb_jargon_group_manage",    "GET",  "/kb/jargon_groups",                     YES, YES, NO,  NO,  NO,  ""),
    ("KB", "kb_jargon_group_manage",    "GET",  "/kb/jargon_groups/jargons",             YES, YES, NO,  NO,  NO,  ""),
    ("KB", "kb_jargon_group_manage",    "GET",  "/kb/jargon_groups/version/{id}",        YES, YES, NO,  NO,  NO,  ""),
    ("KB", "kb_jargon_group_manage",    "POST", "/kb/jargon_groups/add",                 YES, YES, NO,  NO,  NO,  ""),
    ("KB", "kb_jargon_group_manage",    "POST", "/kb/jargon_groups/remove",              YES, YES, NO,  NO,  NO,  ""),
    ("KB", "kb_jargon_group_manage",    "POST", "/kb/jargon_groups/knowledge_bases/add",    YES, YES, NO, NO, NO, ""),
    ("KB", "kb_jargon_group_manage",    "POST", "/kb/jargon_groups/knowledge_bases/remove", YES, YES, NO, NO, NO, ""),
    ("KB", "kb_jargon_manage",          "POST", "/kb/jargons/add",                       YES, YES, NO,  NO,  NO,  ""),
    ("KB", "kb_jargon_manage",          "POST", "/kb/jargons/modify",                    YES, YES, NO,  NO,  NO,  ""),
    ("KB", "kb_jargon_manage",          "POST", "/kb/jargons/remove",                    YES, YES, NO,  NO,  NO,  ""),

    # ========= Rubik =========
    ("Rubik", "rubik_db_import",          "POST",   "/rubik/api/databases",                  YES, NO, YES, NO,  YES, "创建→owner"),
    ("Rubik", "rubik_db_import",          "DELETE", "/rubik/api/databases/{id}",             YES, NO, YES, NO,  RES, "owner"),
    ("Rubik", "rubik_db_view",            "GET",    "/rubik/api/databases",                  YES, NO, YES, NO,  RES, "X-Allowed-Ids"),
    ("Rubik", "rubik_db_view",            "GET",    "/rubik/api/databases/{id}",             YES, NO, YES, NO,  RES, "viewer"),
    ("Rubik", "rubik_db_view",            "GET",    "/rubik/api/databases/{id}/check",       YES, NO, YES, NO,  RES, ""),
    ("Rubik", "rubik_db_view",            "GET",    "/rubik/api/databases/{id}/schema",      YES, NO, YES, NO,  RES, ""),
    ("Rubik", "rubik_db_view",            "GET",    "/rubik/api/databases/{id}/tables-columns", YES, NO, YES, NO, RES, ""),
    ("Rubik", "rubik_db_view",            "GET",    "/rubik/api/databases/{id}/tables/{name}", YES, NO, YES, NO, RES, ""),
    ("Rubik", "rubik_db_view",            "POST",   "/rubik/api/databases/{id}/prettify-sql", YES, NO, YES, NO, RES, "viewer"),
    ("Rubik", "rubik_db_view",            "POST",   "/rubik/api/databases/{id}/execute-sql",  YES, NO, YES, NO, RES, "contributor"),
    ("Rubik", "rubik_db_view",            "POST",   "/rubik/api/databases/test",              YES, NO, YES, NO, YES, ""),
    ("Rubik", "rubik_db_data_dir",        "GET",    "/rubik/api/databases/config/data-dir",   YES, NO, YES, NO, "⚠", "理想 rubik-admins only；OR 语义下被 rubik_db_view GET 覆盖（已知限制）"),
    ("Rubik", "rubik_kb_build",           "POST",   "/rubik/api/databases/{id}/build",        YES, NO, YES, NO, RES, "contributor"),
    ("Rubik", "rubik_kb_build",           "POST",   "/rubik/api/databases/{id}/build/stream", YES, NO, YES, NO, RES, "contributor"),
    ("Rubik", "rubik_kb_build",           "POST",   "/rubik/api/databases/{id}/build/cancel", YES, NO, YES, NO, RES, "contributor"),
    ("Rubik", "rubik_kb_build",           "GET",    "/rubik/api/databases/{id}/build/status", YES, NO, YES, NO, RES, "viewer"),
    ("Rubik", "rubik_knowledge_edit",     "PUT",    "/rubik/api/databases/{id}/knowledge/{type}/{item}", YES, NO, YES, NO, RES, "contributor"),
    ("Rubik", "rubik_knowledge_edit",     "DELETE", "/rubik/api/databases/{id}/knowledge/{type}/{item}", YES, NO, YES, NO, RES, "contributor"),
    ("Rubik", "rubik_knowledge_edit",     "POST",   "/rubik/api/databases/{id}/knowledge/taxonomy",   YES, NO, YES, NO, RES, ""),
    ("Rubik", "rubik_knowledge_edit",     "POST",   "/rubik/api/databases/{id}/knowledge/custom",     YES, NO, YES, NO, RES, ""),
    ("Rubik", "rubik_knowledge_edit",     "POST",   "/rubik/api/databases/{id}/knowledge/experience", YES, NO, YES, NO, RES, ""),
    ("Rubik", "rubik_knowledge_edit",     "POST",   "/rubik/api/databases/{id}/knowledge/import",     YES, NO, YES, NO, RES, ""),
    ("Rubik", "rubik_knowledge_special",  "POST",   "/rubik/api/databases/knowledge/special", YES, NO, YES, NO, "⚠", "理想 rubik-admins only；OR 语义下被 rubik_db_view/edit POST 覆盖（已知限制）"),
    ("Rubik", "rubik_knowledge_view",     "GET",    "/rubik/api/databases/{id}/knowledge/types",      YES, NO, YES, NO, RES, "viewer"),
    ("Rubik", "rubik_knowledge_view",     "GET",    "/rubik/api/databases/{id}/knowledge/list",       YES, NO, YES, NO, RES, ""),
    ("Rubik", "rubik_knowledge_view",     "GET",    "/rubik/api/databases/{id}/knowledge/{type}",     YES, NO, YES, NO, RES, ""),
    ("Rubik", "rubik_knowledge_view",     "GET",    "/rubik/api/databases/{id}/knowledge/{type}/{item}", YES, NO, YES, NO, RES, ""),
    ("Rubik", "rubik_knowledge_view",     "POST",   "/rubik/api/databases/{id}/sync",                 YES, NO, YES, NO, RES, "contributor"),
    ("Rubik", "rubik_knowledge_view",     "POST",   "/rubik/api/databases/{id}/sync/stream",          YES, NO, YES, NO, RES, ""),
    ("Rubik", "rubik_knowledge_view",     "POST",   "/rubik/api/databases/{id}/knowledge/export/stream", YES, NO, YES, NO, RES, "viewer"),
    ("Rubik", "rubik_skill_manage",       "POST",   "/rubik/api/databases/{id}/skill/custom",  YES, NO, YES, NO, RES, "contributor"),
    ("Rubik", "rubik_skill_manage",       "PUT",    "/rubik/api/databases/{id}/skill/{item}",  YES, NO, YES, NO, RES, "contributor"),
    ("Rubik", "rubik_metadata",           "GET",    "/rubik/api/metadata/{id}/init",           YES, NO, YES, NO, RES, "contributor (实为写)"),
    ("Rubik", "rubik_metadata",           "GET",    "/rubik/api/metadata/{id}",                YES, NO, YES, NO, RES, "viewer"),
    ("Rubik", "rubik_metadata",           "PUT",    "/rubik/api/metadata/{id}/description",    YES, NO, YES, NO, RES, "contributor"),
    ("Rubik", "rubik_query",              "POST",   "/rubik/api/query",                        YES, NO, YES, NO, RES, "viewer on body.database_id"),
    ("Rubik", "rubik_query",              "GET",    "/rubik/api/sessions",                     YES, NO, YES, NO, RES, "X-Allowed-Ids"),
    ("Rubik", "rubik_query",              "POST",   "/rubik/api/sessions",                     YES, NO, YES, NO, YES, "创建→owner"),
    ("Rubik", "rubik_query",              "DELETE", "/rubik/api/sessions/{id}",                YES, NO, YES, NO, RES, "owner"),
    ("Rubik", "rubik_query",              "POST",   "/rubik/api/sessions/replay",              YES, NO, YES, NO, RES, ""),
    ("Rubik", "rubik_query",              "GET",    "/rubik/api/sessions/{id}/replay",         YES, NO, YES, NO, RES, "owner"),
    ("Rubik", "rubik_query",              "GET",    "/rubik/api/sessions/{id}/turns",          YES, NO, YES, NO, RES, "owner"),
    ("Rubik", "rubik_query",              "POST",   "/rubik/api/sessions/{id}/turns/{tid}/feedback", YES, NO, YES, NO, RES, "owner"),
    ("Rubik", "rubik_config_view",        "GET",    "/rubik/api/config",                       YES, NO, YES, NO, YES, ""),
    ("Rubik", "rubik_config_view",        "GET",    "/rubik/api/config/models",                YES, NO, YES, NO, YES, ""),
    ("Rubik", "rubik_config_view",        "GET",    "/rubik/api/config/database-providers",    YES, NO, YES, NO, YES, ""),
    ("Rubik", "rubik_config_view",        "GET",    "/rubik/api/config/language",              YES, NO, YES, NO, YES, ""),
    ("Rubik", "rubik_config_view",        "GET",    "/rubik/api/config/app",                   YES, NO, YES, NO, YES, ""),
    ("Rubik", "rubik_config_view",        "GET",    "/rubik/api/config/paths/logs",            YES, NO, YES, NO, YES, ""),
    ("Rubik", "rubik_config_view",        "GET",    "/rubik/api/config/llm-providers",         YES, NO, YES, NO, YES, ""),
    ("Rubik", "rubik_config_manage",      "PUT",    "/rubik/api/config/models/{name}",         YES, NO, YES, NO, NO, "rubik-admins 专属"),
    ("Rubik", "rubik_config_manage",      "PUT",    "/rubik/api/config/database-providers/{p}", YES, NO, YES, NO, NO, ""),
    ("Rubik", "rubik_config_manage",      "PUT",    "/rubik/api/config/language",              YES, NO, YES, NO, NO, ""),
    ("Rubik", "rubik_config_manage",      "PUT",    "/rubik/api/config/languages",             YES, NO, YES, NO, NO, ""),
    ("Rubik", "rubik_config_manage",      "PUT",    "/rubik/api/config/app/{key}",             YES, NO, YES, NO, NO, ""),
    ("Rubik", "rubik_config_manage",      "POST",   "/rubik/api/config/reload",                YES, NO, YES, NO, NO, ""),
    ("Rubik", "rubik_config_manage",      "POST",   "/rubik/api/config/setup",                 YES, NO, YES, NO, NO, ""),
    ("Rubik", "rubik_config_manage",      "POST",   "/rubik/api/config/open-path",             YES, NO, YES, NO, NO, ""),
    ("Rubik", "rubik_config_manage",      "POST",   "/rubik/api/config/llm-providers",         YES, NO, YES, NO, NO, ""),
    ("Rubik", "rubik_config_manage",      "PUT",    "/rubik/api/config/llm-providers/{name}",  YES, NO, YES, NO, NO, ""),
    ("Rubik", "rubik_config_manage",      "DELETE", "/rubik/api/config/llm-providers/{name}",  YES, NO, YES, NO, NO, ""),
    ("Rubik", "rubik_dashboard_view",     "GET",    "/rubik/api/dashboards",                   YES, NO, YES, NO, YES, ""),
    ("Rubik", "rubik_dashboard_view",     "GET",    "/rubik/api/dashboards/{id}",              YES, NO, YES, NO, YES, ""),
    ("Rubik", "rubik_dashboard_edit",     "POST",   "/rubik/api/dashboards",                   YES, NO, YES, NO, YES, ""),
    ("Rubik", "rubik_dashboard_edit",     "PUT",    "/rubik/api/dashboards/{id}",              YES, NO, YES, NO, YES, ""),
    ("Rubik", "rubik_dashboard_edit",     "DELETE", "/rubik/api/dashboards/{id}",              YES, NO, YES, NO, YES, ""),
    ("Rubik", "rubik_db_refresh",         "GET",    "/rubik/api/refresh/{id}/status",          YES, NO, YES, NO, YES, ""),
    ("Rubik", "rubik_db_refresh",         "POST",   "/rubik/api/refresh/{id}/execute/stream",  YES, NO, YES, NO, YES, ""),

    # ========= Memory =========
    ("Memory", "memory_health",            "GET",    "/memory/api/v1/health",             YES, NO, NO, YES, YES, ""),
    ("Memory", "memory_system",            "POST",   "/memory/api/v1/system/recovery",    YES, NO, NO, YES, YES, "semantically admin; 按用户意图=all-users"),
    ("Memory", "memory_data_rw",           "POST",   "/memory/api/v1/memory/add",         YES, NO, NO, YES, YES, ""),
    ("Memory", "memory_data_rw",           "POST",   "/memory/api/v1/memory/query",       YES, NO, NO, YES, YES, ""),
    ("Memory", "memory_data_rw",           "POST",   "/memory/api/v1/memory/update",      YES, NO, NO, YES, YES, ""),
    ("Memory", "memory_data_rw",           "POST",   "/memory/api/v1/memory/delete",      YES, NO, NO, YES, YES, ""),
    ("Memory", "memory_tenant_manage",     "POST",   "/memory/api/v1/tenants",            YES, NO, NO, YES, NO,  "管理面仅 memory-admins"),
    ("Memory", "memory_tenant_manage",     "DELETE", "/memory/api/v1/tenants/{id}",       YES, NO, NO, YES, NO,  ""),
    ("Memory", "memory_tenant_manage",     "POST",   "/memory/api/v1/tenants/{id}/instances",          YES, NO, NO, YES, NO, ""),
    ("Memory", "memory_tenant_manage",     "DELETE", "/memory/api/v1/tenants/{id}/instances/{name}",   YES, NO, NO, YES, NO, ""),
    ("Memory", "memory_tenant_manage",     "DELETE", "/memory/api/v1/tenants/{id}/instances/{n}/users/{u}/memories", YES, NO, NO, YES, NO, ""),
    ("Memory", "memory_template_manage",   "POST",   "/memory/api/v1/templates",          YES, NO, NO, YES, NO, ""),
    ("Memory", "memory_template_manage",   "GET",    "/memory/api/v1/templates",          YES, NO, NO, YES, NO, ""),
    ("Memory", "memory_template_manage",   "GET",    "/memory/api/v1/templates/{id}",     YES, NO, NO, YES, NO, ""),
    ("Memory", "memory_template_manage",   "PUT",    "/memory/api/v1/templates/{id}",     YES, NO, NO, YES, NO, ""),
    ("Memory", "memory_template_manage",   "DELETE", "/memory/api/v1/templates/{id}",     YES, NO, NO, YES, NO, ""),
    ("Memory", "memory_template_manage",   "POST",   "/memory/api/v1/templates/{id}/filters",         YES, NO, NO, YES, NO, ""),
    ("Memory", "memory_template_manage",   "GET",    "/memory/api/v1/templates/{id}/filters",         YES, NO, NO, YES, NO, ""),
    ("Memory", "memory_template_manage",   "POST",   "/memory/api/v1/templates/{id}/llm-extraction",  YES, NO, NO, YES, NO, ""),
    ("Memory", "memory_template_manage",   "POST",   "/memory/api/v1/templates/batch/llm-extraction", YES, NO, NO, YES, NO, ""),
]


def _header_at(ws, row, headers):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = hf; cell.fill = hfill; cell.alignment = center; cell.border = border


def build():
    wb = Workbook()

    # ══════════════════════════════════════════════════════════════════════
    # Sheet 1: IAM 系统接口
    # ══════════════════════════════════════════════════════════════════════
    ws1 = wb.active
    ws1.title = "1. IAM系统接口"
    _header(ws1, IAM_HEADERS)
    mod_fill = {
        "用户管理":   PatternFill("solid", fgColor="D9E2F3"),
        "用户组管理": PatternFill("solid", fgColor="E2EFDA"),
        "权限管理":   PatternFill("solid", fgColor="FFF2CC"),
        "应用管理":   PatternFill("solid", fgColor="FCE4D6"),
        "API Key":    PatternFill("solid", fgColor="F2DCDB"),
        "资源ACL":    PatternFill("solid", fgColor="E4DFEC"),
        "IdP管理":    PatternFill("solid", fgColor="D5E8D4"),
        "Token":      PatternFill("solid", fgColor="DAE8FC"),
    }
    for r, api in enumerate(IAM_APIS, 2):
        _row(ws1, r, api, method_col=3)
        ws1.cell(row=r, column=1).fill = mod_fill.get(api[0], PatternFill())
    for col, w in {1: 12, 2: 22, 3: 8, 4: 54, 5: 14, 6: 48, 7: 42, 8: 48, 9: 28}.items():
        ws1.column_dimensions[get_column_letter(col)].width = w
    ws1.row_dimensions[1].height = 24
    for r in range(2, len(IAM_APIS) + 2):
        ws1.row_dimensions[r].height = 70
    ws1.freeze_panes = "E2"
    ws1.auto_filter.ref = f"A1:I{len(IAM_APIS) + 1}"

    # ══════════════════════════════════════════════════════════════════════
    # Sheet 2: 权限点总表
    # ══════════════════════════════════════════════════════════════════════
    ws2 = wb.create_sheet("2. 权限点总表")
    _header(ws2, PG_HEADERS)
    app_fill = {
        "":              PatternFill("solid", fgColor="FFE699"),
        "knowledgebase": PatternFill("solid", fgColor="D9E2F3"),
        "rubik":         PatternFill("solid", fgColor="E2EFDA"),
        "memory":        PatternFill("solid", fgColor="FCE4D6"),
    }
    row = 2
    pg_id = 1
    for app_name, name, desc, paths, kc_groups in PERM_GROUPS:
        # Expand (path, method) pairs into multiple rows sharing the same group id
        for i, (path, method) in enumerate(paths):
            values = (
                pg_id if i == 0 else "",
                app_name if app_name else "(平台级)",
                name if i == 0 else "",
                desc if i == 0 else "",
                path,
                method,
                ", ".join(kc_groups) if i == 0 else "",
            )
            for c, v in enumerate(values, 1):
                cell = ws2.cell(row=row, column=c, value=v)
                cell.font = bf
                cell.alignment = wrap if c >= 4 else center
                cell.border = border
                if c == 6:  # method
                    cell.fill = method_fill.get(method, PatternFill())
                if c == 2:
                    cell.fill = app_fill.get(app_name, PatternFill())
            row += 1
        pg_id += 1
    for col, w in {1: 6, 2: 15, 3: 28, 4: 44, 5: 44, 6: 9, 7: 28}.items():
        ws2.column_dimensions[get_column_letter(col)].width = w
    ws2.row_dimensions[1].height = 24
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:G{row - 1}"

    # ══════════════════════════════════════════════════════════════════════
    # Sheet 3: 群组 × 应用 概览
    # ══════════════════════════════════════════════════════════════════════
    ws3 = wb.create_sheet("3. 群组×应用概览")
    _header(ws3, OVERVIEW_HEADERS)
    for r, row_vals in enumerate(OVERVIEW_ROWS, 2):
        for c, v in enumerate(row_vals, 1):
            cell = ws3.cell(row=r, column=c, value=v)
            cell.font = bf; cell.alignment = wrap if c <= 2 else center; cell.border = border
            if c == 1:
                cell.font = bf_bold
                cell.fill = PatternFill("solid", fgColor="FFE699")
            elif c >= 3:
                if "✓" in v:
                    cell.fill = PatternFill("solid", fgColor="C6EFCE")
                elif "⚡" in v:
                    cell.fill = PatternFill("solid", fgColor="FFEB9C")
                elif v == "✗":
                    cell.fill = PatternFill("solid", fgColor="FFC7CE")
    for col, w in {1: 16, 2: 26, 3: 14, 4: 28, 5: 38, 6: 38, 7: 28}.items():
        ws3.column_dimensions[get_column_letter(col)].width = w
    ws3.row_dimensions[1].height = 24
    for r in range(2, len(OVERVIEW_ROWS) + 2):
        ws3.row_dimensions[r].height = 40

    # Legend below
    legend_row = len(OVERVIEW_ROWS) + 4
    ws3.cell(row=legend_row, column=1, value="图例").font = bf_bold
    for i, (sym, desc) in enumerate([
        ("✓ 全权", "通过 {app}_admin_full 或具体 permission_group 放行，无资源级限制"),
        ("⚡ owner/contrib/viewer", "路径级放行，资源级 ACL 按 permission 梯度再过滤"),
        ("✗", "路径级默认拒绝（Default Deny，无匹配 permission_group）"),
    ]):
        cell = ws3.cell(row=legend_row + 1 + i, column=1, value=sym)
        cell.font = bf_bold
        ws3.cell(row=legend_row + 1 + i, column=2, value=desc).font = bf

    # ══════════════════════════════════════════════════════════════════════
    # Sheet 4: 业务接口权限矩阵
    # ══════════════════════════════════════════════════════════════════════
    ws4 = wb.create_sheet("4. 业务接口权限矩阵")
    _header(ws4, ENDPOINT_HEADERS)
    app_band = {"KB": "D9E2F3", "Rubik": "E2EFDA", "Memory": "FCE4D6"}
    for r, row_vals in enumerate(ENDPOINTS, 2):
        app, pg, method, path, *perm_cells, note = row_vals[:9] + (row_vals[9],)
        # values already 10-tuple; _row + manual styling
        for c, v in enumerate(row_vals, 1):
            cell = ws4.cell(row=r, column=c, value=v)
            cell.font = bf
            cell.alignment = center if c in (1, 3, 5, 6, 7, 8, 9) else wrap
            cell.border = border
            if c == 1:
                cell.fill = PatternFill("solid", fgColor=app_band.get(v, "FFFFFF"))
                cell.font = bf_bold
            elif c == 3:  # method
                cell.fill = method_fill.get(v, PatternFill())
            elif c in (5, 6, 7, 8, 9):  # perm cells
                if v == YES:
                    cell.fill = PatternFill("solid", fgColor="C6EFCE")
                elif v == RES:
                    cell.fill = PatternFill("solid", fgColor="FFEB9C")
                elif v == NO:
                    cell.fill = PatternFill("solid", fgColor="FFC7CE")
                elif v == "⚠":
                    cell.fill = PatternFill("solid", fgColor="F4B084")
    for col, w in {1: 8, 2: 28, 3: 9, 4: 54, 5: 9, 6: 12, 7: 14, 8: 14, 9: 10, 10: 40}.items():
        ws4.column_dimensions[get_column_letter(col)].width = w
    ws4.row_dimensions[1].height = 28
    ws4.freeze_panes = "D2"
    ws4.auto_filter.ref = f"A1:J{len(ENDPOINTS) + 1}"

    # ---- Save ----
    out = "diagrams/api-spec-v2.xlsx"
    wb.save(out)
    print(
        f"OK: {out}\n"
        f"  sheet 1 (IAM APIs): {len(IAM_APIS)} rows\n"
        f"  sheet 2 (permission_groups): {len(PERM_GROUPS)} groups\n"
        f"  sheet 3 (group x app overview): {len(OVERVIEW_ROWS)} groups\n"
        f"  sheet 4 (endpoint x group matrix): {len(ENDPOINTS)} endpoints"
    )


if __name__ == "__main__":
    build()
